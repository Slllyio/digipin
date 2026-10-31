"""Homes and households in 2011 by ward (Mumbai; Census 2011 ward tables via OpenCity, public domain).

The 2011 Census split Greater Mumbai into 97 census wards, each inside one of BMC's 24 administrative wards (A to T).
OpenCity publishes the Primary Census Abstract and the Houselisting and Housing table HH-14 (percentage of households by
amenities and assets) per census ward. The census wards are summed into the administrative wards, weighting HH-14's
percentages by households, and every L7 block in a BMC ward (``wards`` dataset) carries its ward's figures
(dataset ``census_homes``):

    c11_ward             BMC administrative ward
    c11_ward_no          its position in BMC_WARDS (A = 1 ... T = 24), for lenses that read numbers only
    c11_households       households
    c11_hh_size          people per household
    c11_literacy_pct     literate people among those aged 7 or more (%)
    c11_women_work_pct   women who work (main or marginal) (%)
    c11_dilapidated_pct  households in a dilapidated house (%)
    c11_one_room_pct     households in one room or with no room of their own (%)
    c11_rented_pct       households renting (%)
    c11_tap_pct          households drinking treated tap water (%)
    c11_water_away_pct   households fetching drinking water from away (%)
    c11_no_latrine_pct   households with no latrine of their own (%)
    c11_open_def_pct     households with no latrine that go in the open (%)
    c11_closed_drain_pct households whose waste water goes to a closed drain (%)
    c11_lpg_pct          households cooking on LPG or piped gas (%)
    c11_bank_pct         households with a bank account (%)
    c11_internet_pct     households with a computer and internet (%)
    c11_car_pct          households with a car, jeep or van (%)
    c11_twowheeler_pct   households with a scooter or motorcycle (%)
    c11_no_assets_pct    households with none of the listed assets (%)

Caste and tribe columns are never read.

    python -m pipeline.realestate.census2011 [--region mumbai]
"""
from __future__ import annotations

import argparse
import csv
import io
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
CELLS = REPO / "data" / "cells"
LEVEL = 7
PACKAGE = "mumbai-ward-wise-census-data"
# HH-14 column (0-based) → field; several columns are summed.
HH14 = {
    "c11_dilapidated_pct": (13,), "c11_one_room_pct": (48, 49), "c11_rented_pct": (63,), "c11_tap_pct": (71,),
    "c11_water_away_pct": (83,), "c11_no_latrine_pct": (99,), "c11_open_def_pct": (101,), "c11_closed_drain_pct": (105,),
    "c11_lpg_pct": (113,), "c11_bank_pct": (126,), "c11_internet_pct": (129,), "c11_car_pct": (136,),
    "c11_twowheeler_pct": (135,), "c11_no_assets_pct": (138,),
}
BMC_WARDS = ("A", "B", "C", "D", "E", "F/N", "F/S", "G/N", "G/S", "H/E", "H/W", "K/E", "K/W", "L", "M/E", "M/W", "N", "P/N",
             "P/S", "R/C", "R/N", "R/S", "S", "T")
PCA = ("Ward", "Level", "TRU", "No_HH", "TOT_P", "TOT_F", "P_06", "P_LIT", "TOT_WORK_F")


def _f(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def admin_wards(rows):
    """{census ward code: admin ward} from the census-ward list (Ward Code, Ward Name)."""
    return {int(r["Ward Code"]): r["Ward Name"].strip() for r in rows if (r.get("Ward Code") or "").strip().isdigit()}


def pca(rows):
    """{census ward code: counts} from Primary Census Abstract rows at ward level (Total rows)."""
    out = {}
    for r in rows:
        if r.get("Level") != "WARD" or r.get("TRU") not in ("Total", "Urban"):
            continue
        out[int(r["Ward"])] = {k: _f(r.get(k)) or 0 for k in ("No_HH", "TOT_P", "TOT_F", "P_06", "P_LIT", "TOT_WORK_F")}
    return out


def hh14(rows):
    """{census ward code: {field: %}} from HH-14 sheet rows (tuples); ward rows have a non-zero ward number."""
    out = {}
    for r in rows:
        if len(r) < 139 or not str(r[7] or "").strip().isdigit() or int(r[7]) == 0 or r[9] not in ("Total", "Urban"):
            continue
        vals = {}
        for k, cols in HH14.items():
            xs = [_f(r[c]) for c in cols]
            if all(x is not None for x in xs):
                vals[k] = sum(xs)
        out.setdefault(int(r[7]), vals)
    return out


def wards(names, counts, shares):
    """{admin ward: figures}: counts summed, HH-14 shares weighted by households."""
    acc = {}
    for code, c in counts.items():
        w = names.get(code)
        if not w:
            continue
        a = acc.setdefault(w, {"hh": 0, "p": 0, "f": 0, "p06": 0, "lit": 0, "fw": 0, "share": {}, "wt": {}})
        a["hh"] += c["No_HH"]
        a["p"] += c["TOT_P"]
        a["f"] += c["TOT_F"]
        a["p06"] += c["P_06"]
        a["lit"] += c["P_LIT"]
        a["fw"] += c["TOT_WORK_F"]
        for k, v in shares.get(code, {}).items():
            a["share"][k] = a["share"].get(k, 0) + v * c["No_HH"]
            a["wt"][k] = a["wt"].get(k, 0) + c["No_HH"]
    out = {}
    for w, a in acc.items():
        if not a["hh"]:
            continue
        out[w] = {"c11_households": int(a["hh"]), "c11_hh_size": round(a["p"] / a["hh"], 2),
                  **({"c11_literacy_pct": round(100 * a["lit"] / (a["p"] - a["p06"]), 1)} if a["p"] > a["p06"] else {}),
                  **({"c11_women_work_pct": round(100 * a["fw"] / a["f"], 1)} if a["f"] else {}),
                  **{k: round(a["share"][k] / a["wt"][k], 1) for k in a["share"] if a["wt"][k]}}
    return out


def ward_cells(region, root=CELLS, body="BMC"):
    """[(code, admin ward)] for blocks of ``body`` in the wards dataset."""
    import json

    from pipeline.cells import store
    cat = json.loads((root / "catalog.json").read_text())
    ds = next(d for d in cat["datasets"] if d["id"] == "wards")
    t = store.read_region(root, next(r for r in ds["regions"] if r["name"] == region))
    return [(c, w) for c, w, b in zip(t.codes, t.columns["ward"], t.columns["body"]) if w and b == body]


def records(cells, ws):
    for code, w in cells:
        if w in ws:
            yield {"digipin": code, "c11_ward": w, **({"c11_ward_no": BMC_WARDS.index(w) + 1} if w in BMC_WARDS else {}), **ws[w]}


LABELS = {
    "c11_households": ("Households in the ward, 2011", "households", None, None),
    "c11_hh_size": ("People per household, 2011", "people", [3, 6], None),
    "c11_literacy_pct": ("Literate people aged 7 or more, 2011", "%", [70, 100], "better"),
    "c11_women_work_pct": ("Women who work, 2011", "%", [0, 30], None),
    "c11_dilapidated_pct": ("Households in a dilapidated house, 2011", "%", [0, 5], "worse"),
    "c11_one_room_pct": ("Households in one room or less, 2011", "%", [0, 80], "worse"),
    "c11_rented_pct": ("Households renting, 2011", "%", [0, 60], None),
    "c11_tap_pct": ("Households drinking treated tap water, 2011", "%", [70, 100], "better"),
    "c11_water_away_pct": ("Households fetching water from away, 2011", "%", [0, 10], "worse"),
    "c11_no_latrine_pct": ("Households without a latrine of their own, 2011", "%", [0, 70], "worse"),
    "c11_open_def_pct": ("Households going in the open, 2011", "%", [0, 5], "worse"),
    "c11_closed_drain_pct": ("Households on a closed drain, 2011", "%", [30, 100], "better"),
    "c11_lpg_pct": ("Households cooking on LPG or piped gas, 2011", "%", [50, 100], "better"),
    "c11_bank_pct": ("Households with a bank account, 2011", "%", [50, 100], "better"),
    "c11_internet_pct": ("Households with a computer and internet, 2011", "%", [0, 40], None),
    "c11_car_pct": ("Households with a car, 2011", "%", [0, 40], None),
    "c11_twowheeler_pct": ("Households with a scooter or motorcycle, 2011", "%", [0, 40], None),
    "c11_no_assets_pct": ("Households with none of the listed assets, 2011", "%", [0, 10], "worse"),
}


def cell_dataset(region, recs):
    from pipeline.cells import store
    from pipeline.cells.schema import DatasetSpec, FieldSpec
    specs = [FieldSpec("c11_ward", "BMC ward", type="category", agg="mode"),
             FieldSpec("c11_ward_no", "BMC ward number (A = 1 ... T = 24)", agg="min", ramp="viridis", higher_is=None)]
    for k, (label, unit, dom, hi) in LABELS.items():
        specs.append(FieldSpec(k, label, agg="max" if k == "c11_households" else "mean", unit=unit, ramp="viridis",
                               higher_is=hi, **({"domain": dom} if dom else {})))
    spec = DatasetSpec(
        id="census_homes", title="Homes and households in 2011 by ward (Census 2011)", level=LEVEL, fields=specs, group="Real estate",
        description="Census 2011 Primary Census Abstract and Houselisting table HH-14 for Greater Mumbai's 97 census wards on "
                    "data.opencity.in, summed into BMC's 24 administrative wards (shares weighted by households): crowding, "
                    "renting, water, sanitation, cooking fuel, banking and household assets. Every block in a ward carries the "
                    "ward's figures.",
        source="pipeline/realestate/census2011.py (OpenCity Mumbai ward-wise census data)",
        license="Public domain (Census of India 2011 via data.opencity.in)", default_field="c11_one_room_pct", extrude=False)
    return store.write_dataset(CELLS, spec, region, store.bin_records(recs, LEVEL, specs), parquet=False)


def main(argv=None):
    import openpyxl
    import requests
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--region", default="mumbai")
    region = ap.parse_args(argv).region
    res = requests.get("https://data.opencity.in/api/3/action/package_show", params={"id": PACKAGE}, timeout=60).json()["result"]["resources"]
    names, counts, shares = {}, {}, {}
    for r in res:
        body = requests.get(r["url"], timeout=300).content
        if r["format"] == "CSV":
            rows = list(csv.DictReader(io.StringIO(body.decode("utf-8-sig"))))
            if rows and "Ward Code" in rows[0]:
                names.update(admin_wards(rows))
            elif rows and "Level" in rows[0]:
                counts.update(pca({k: row.get(k) for k in PCA} for row in rows))
        elif r["format"] == "XLSX":
            ws = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True).worksheets[0]
            shares.update(hh14(ws.iter_rows(values_only=True)))
    ws = wards(names, counts, shares)
    info = cell_dataset(region, list(records(ward_cells(region), ws)))
    print(f"{region}: {len(counts)} census wards, {len(shares)} with HH-14, {len(ws)} BMC wards, {info['cells']} blocks")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
